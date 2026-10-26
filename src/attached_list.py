import sys
from typing import List, Dict, Any, Optional
import platform
import openpyxl
from openpyxl.styles.borders import Border, Side
from IUpdate_master import IUpdateMaster
from get_idx import GetIdx

class AttachedList(IUpdateMaster):
    def __init__(self, hinban: str, hin_name: str,
                       coaFormat_col: List[str], # 向先_ｼｰﾄ表のデータ
                       coaFormat_data: List[List[str]], # 向先_ｼｰﾄ表のデータ
                )-> None:

        self._hinban = hinban
        self._hin_name = hin_name
        self._coaFormat_col = coaFormat_col
        self._coaFormat_data = coaFormat_data

        self._path: str = r'\\192.168.1.247\共有\営業課ﾌｫﾙﾀﾞ\櫻田\☆☆☆\売上処理(水野課長用)\出荷時添付リスト(20200731時点最新版).xlsx'
        if platform.system() == 'Linux':
            self._path = r'/mnt/public/営業課ﾌｫﾙﾀﾞ/櫻田/☆☆☆/売上処理(水野課長用)/出荷時添付リスト(20200731時点最新版).xlsx'

        # Excelブックを開く
        try:
            self._wb = openpyxl.load_workbook(self._path, data_only=True)
            self._ws = self._wb['成績表']
        except Exception as e:
            print(f'{e} 出荷時添付リストを開けません。プログラムを終了します。')
            sys.exit(1)

        # 全ての行を取得
        rows = list(self._ws.iter_rows(values_only=True))
        # 2行目をカラム名として取得
        self._columns = [str(col) for col in rows[1]]
        # 3行目以降をデータとして取得
        self._data_list = [list(row) for row in rows[2:]]


    def _slice_nonyuCD(self, strs: str)-> str:
        # strs = 'H0013   -> 横浜真空' など
        nonyuCD: str = ''
        strs_list = strs.split()
        if len(strs_list) == 3:
            return strs_list[0]
        if len(strs_list) == 2:
            return ''
        return nonyuCD


    def _slice_tokuisaki_name(self, strs: str)-> str:
        # strs = '    -> 横浜真空' など
        strs_list = strs.split()
        return strs_list[-1]


    def _get_atena(self, coaFormat: str)-> str:
        atena : str = '-'
        coaFormat_idx:int = GetIdx.get_idx(self._coaFormat_col, '向先')
        atena_idx: int = GetIdx.get_idx(self._coaFormat_col, '宛名')
        for row in self._coaFormat_data:
            if row[coaFormat_idx] == coaFormat:
                return row[atena_idx]
        return atena


    def _check_is_exist_data(self, tokuiCD: str, nonyuCD: Optional[str],
                            ) -> bool:

        tokuiCD_idx:int = GetIdx.get_idx(self._columns, '得意先ｺｰﾄﾞ')
        nonyuCD_idx:int = GetIdx.get_idx(self._columns, '納入先ｺｰﾄﾞ')
        hinban_idx:int = GetIdx.get_idx(self._columns, '品番')
        for row in self._data_list:
            if (row[tokuiCD_idx] == tokuiCD and 
                row[nonyuCD_idx] == nonyuCD and 
                row[hinban_idx] == self._hinban):
                return True
        return False
        

    def _update_data(self, tokuiCD: str, nonyuCD: Optional[str], 
                     coaFormat: str) -> None:
        tokuiCD_idx:int = GetIdx.get_idx(self._columns, '得意先ｺｰﾄﾞ')
        nonyuCD_idx:int = GetIdx.get_idx(self._columns, '納入先ｺｰﾄﾞ')
        hinban_idx:int = GetIdx.get_idx(self._columns, '品番')
        coa_format_idx:int = GetIdx.get_idx(self._columns, 'format')
        department_idx:int = GetIdx.get_idx(self._columns, '発行部署')
        department = '営業'
        if coaFormat == '小糸':
            department = '技術'


        if nonyuCD == '':
            nonyuCD = None
        for row in range(3, self._ws.max_row + 1):
            if (self._ws.cell(row, tokuiCD_idx +1).value == tokuiCD and 
                self._ws.cell(row, nonyuCD_idx +1).value == nonyuCD and 
                self._ws.cell(row, hinban_idx +1).value == self._hinban): 
                self._ws.cell(row, coa_format_idx + 1).value = \
                                                 coaFormat # type: ignore
                self._ws.cell(row, department_idx + 1).value = \
                                                 department # type: ignore
                break
        self._wb.save(self._path)
        self._wb.close()

        print(f'\n {tokuiCD}:{nonyuCD} の{self._hinban}の添付リストを更新しました。')


    def _insert_data(self, insert_list: List[Any]) -> None:

        # 最終行はA列の下から探す(空白である間はデクリメントする)
        last_row = self._ws.max_row
        while last_row > 0 and self._ws.cell(row=last_row, column=1).value is None:
            last_row -= 1

        thin_side = Side(border_style='thin', color='000000')
        dot_side = Side(border_style='dotted', color='000000')
        left_dot = Border(left=dot_side, bottom=thin_side)
        left_thin = Border(left=thin_side, bottom=thin_side)
        left_dot_right_thin = Border(left=dot_side, bottom=thin_side, right=thin_side)
        self._ws.cell(last_row + 1, 1).value = insert_list[0]
        self._ws.cell(last_row + 1, 1).border = left_thin
        self._ws.cell(last_row + 1, 2).value = insert_list[1]
        self._ws.cell(last_row + 1, 2).border = left_dot
        self._ws.cell(last_row + 1, 3).value = insert_list[2]
        self._ws.cell(last_row + 1, 3).border = left_thin
        self._ws.cell(last_row + 1, 4).value = insert_list[3]
        self._ws.cell(last_row + 1, 4).border = left_dot
        self._ws.cell(last_row + 1, 5).value = insert_list[4]
        self._ws.cell(last_row + 1, 5).border = left_dot
        self._ws.cell(last_row + 1, 6).value = insert_list[5]
        self._ws.cell(last_row + 1, 6).border = left_dot
        self._ws.cell(last_row + 1, 7).value = insert_list[6]
        self._ws.cell(last_row + 1, 7).border = left_dot
        self._ws.cell(last_row + 1, 8).value = insert_list[7]
        self._ws.cell(last_row + 1, 8).border = left_dot
        self._ws.cell(last_row + 1, 9).value = insert_list[8]
        self._ws.cell(last_row + 1, 9).border = left_dot_right_thin

        self._wb.save(self._path)
        self._wb.close()

        print(f'\n {insert_list[0]}:{insert_list[1]} の{self._hinban}の添付リストを追加しました。')


    def update_master(self, ui_infos: Dict)-> None:

        tokuiCD:str = ui_infos['tokuiCD']
        nonyuCD:Optional[str] = self._slice_nonyuCD(ui_infos['nonyuCD']) or None
        tokuisaki_name:str = self._slice_tokuisaki_name(ui_infos['nonyuCD'])
        coaFormat:str = ui_infos['format']
        department:str = '営業'
        if coaFormat == '小糸':
            department = '技術'
        atena:str = self._get_atena(coaFormat)
        bikou: Optional[str] = None

        if not ui_infos['is_coa_need']:
            print(f'\n 添付リストの修正、追加は不要です。')
            return

        is_exist_data = self._check_is_exist_data( tokuiCD, nonyuCD)

        if is_exist_data:
            print(f'\n {tokuiCD}:{nonyuCD} の{self._hinban}の添付リストはすでに存在します。')
            self._update_data(
                    tokuiCD, nonyuCD, coaFormat
                )
            return

        if not is_exist_data:
            insert_list: List[Any] = [tokuiCD, nonyuCD, tokuisaki_name, 
                                      self._hinban, self._hin_name, department,
                                      coaFormat, atena, bikou]
            self._insert_data(insert_list)
            return




